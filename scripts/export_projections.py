"""
Export every projected player's stats to one Excel workbook.

    python scripts/export_projections.py                          # -> exports/player_projections_2026-27.xlsx
    python scripts/export_projections.py --out "C:/path/file.xlsx"
    python scripts/export_projections.py --compare old_projections_2026.parquet --compare-label "without team context"   # highlight the biggest changes

Sheets: Projections (per game, one row per player, board order, with games played mean / typical / P25-P75 and the nine z-scores),
Season totals (per game x projected games), Known absences, Injury-impacted (players whose recent games played were hit by injuries),
League GP means, Notes. Run build_projections.py (+ build_library.py) first so the projections are current.
"""
from __future__ import annotations

import argparse
import sqlite3
import sys
from pathlib import Path

import numpy as np
import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.stdout.reconfigure(encoding="utf-8")

from src.draft.values import CAT_LABELS, CATEGORIES, z_scores
from src.features.player_seasons import load_player_season_table

HIST_YEARS = (2022, 2023, 2024, 2025)


def injury_groups(proj: pd.DataFrame, st: pd.DataFrame, rookie_year: dict) -> pd.DataFrame:
    """
    'Injury-hit 2025-26': under 60% of games last season after a real rotation role (>=25 MPG at >=50% of games in 2022-24).
    'Healthy 2025-26, wrecked year before': >=75% of games last season after a <50% season in 2023-24 / 2024-25.
    Kept when per-game rank <= 60 or the board rank at 82% games would be <= 30 (Giannis always). Injury is inferred from games missed.
    """
    s = st.set_index(["player_id", "year"])
    pg_rank = z_scores(proj.assign(avail=1.0))["z_total"].rank(ascending=False, method="first").astype(int)
    out = {}
    for i, p in proj.iterrows():
        h = {}
        for y in HIST_YEARS:
            ry = rookie_year.get(p.player_id)
            if ry is not None and not pd.isna(ry) and y < ry:
                h[y] = (np.inf, np.nan, np.nan)                      # not in the NBA yet: never counts as "wrecked"
            elif (p.player_id, y) in s.index:
                r = s.loc[(p.player_id, y)]
                h[y] = (r["gp"] / r["sg"], r["min"] / r["gp"], int(r["gp"]))
            else:
                h[y] = (0.0, np.nan, 0)
        a = {y: h[y][0] for y in h}; mp = {y: h[y][1] for y in h}
        had_role = any(not np.isnan(mp[y]) and mp[y] >= 25 and a[y] >= 0.5 for y in (2022, 2023, 2024))
        played_before = any(0 < a[y] < np.inf for y in (2022, 2023, 2024))
        hit = a[2025] < 0.60 and had_role
        wrecked = min(a[2023], a[2024]) < 0.50 and a[2025] >= 0.75 and had_role
        if not ((hit or wrecked) and played_before):
            continue
        cf = proj.copy(); cf.loc[i, "avail"] = max(p.avail, 0.82)
        healthy_rank = int(z_scores(cf)["z_total"].rank(ascending=False, method="first").iloc[i])
        if pg_rank[i] <= 60 or healthy_rank <= 30 or p["name"] == "Giannis Antetokounmpo":
            out[p.player_id] = dict(group="Injury-hit 2025-26" if hit else "Healthy 2025-26, wrecked year before",
                                    gp_hist=" / ".join("-" if np.isnan(h[y][2]) else str(int(h[y][2])) for y in HIST_YEARS), rank_if_healthy=healthy_rank)
    return pd.DataFrame.from_dict(out, orient="index")


GREEN, RED, YELLOW = "C6EFCE", "FFC7CE", "FFF2CC"
# a change is "big" (highlighted) at or above these absolute sizes; only players in the top 150 (before or after) are highlighted
BIG = {"z": 1.0, "rank": 10, "pts": 1.5, "mpg": 1.5, "gp": 4.0, "fga": 1.2, "ast": 0.7, "reb": 0.7, "tpm": 0.4}   # z = fantasy value (9-cat z total)


def changes_vs(proj: pd.DataFrame, compare: str, conn: sqlite3.Connection, year: int) -> pd.DataFrame:
    """Per player: previous values from another projections file, deltas, which are big, and whether he changed teams since last season."""
    prev = pd.read_parquet(compare).set_index("player_id")
    p = proj.set_index("player_id")
    c = pd.DataFrame(index=p.index)
    c["prev_rank"] = prev["rank"].reindex(p.index)
    c["d_rank"] = c["prev_rank"] - p["rank"]                                   # + = moved up the board
    c["prev_z"] = prev["z_total"].reindex(p.index); c["d_z"] = p["z_total"] - c["prev_z"]
    for k, col in (("pts", "pts"), ("mpg", "mpg"), ("fga", "fga"), ("ast", "ast"), ("reb", "reb"), ("tpm", "tpm")):
        c[f"prev_{k}"] = prev[col].reindex(p.index); c[f"d_{k}"] = p[col] - c[f"prev_{k}"]
    c["prev_gp"] = prev["avail"].reindex(p.index) * 82; c["d_gp"] = p["avail"] * 82 - c["prev_gp"]
    last = pd.read_sql_query(
        """SELECT bs.player_id, bs.team_id AS last_team, g.date FROM box_score bs JOIN game g ON g.game_id = bs.game_id
           WHERE g.season = ? AND g.game_type = 'regular' AND bs.min > 0""", conn, params=(f"{year - 1}-{str(year)[-2:]}",)
    ).sort_values("date").drop_duplicates("player_id", keep="last").set_index("player_id")["last_team"]
    c["last_team_id"] = last.reindex(p.index)
    c["moved"] = c["last_team_id"].notna() & (c["last_team_id"] != p["team_id"])
    in_top = (p["rank"] <= 150) | (c["prev_rank"] <= 150)
    flags = pd.DataFrame({k: in_top & (c[f"d_{k}"].abs() >= BIG[k]) for k in BIG})
    c["any_big"] = flags.any(axis=1)
    for k in BIG:
        c[f"big_{k}"] = flags[k]
    return c.reset_index()


def main(db_path: str, year: int, out: str, compare: str | None = None, compare_label: str = "previous") -> None:
    from nba_api.stats.static import teams as static_teams

    conn = sqlite3.connect(db_path)
    proj = pd.read_parquet(ROOT / f"data/processed/projections_{year}.parquet").sort_values("rank").reset_index(drop=True)
    st = load_player_season_table(conn)
    rookie_year = pd.read_sql_query("SELECT player_id, rookie_year FROM player", conn).set_index("player_id")["rookie_year"].to_dict()
    abbr = {t["id"]: t["abbreviation"] for t in static_teams.get_teams()}
    inj = injury_groups(proj, st, rookie_year)
    proj["pg_rank"] = z_scores(proj.assign(avail=1.0))["z_total"].rank(ascending=False, method="first").astype(int)
    proj["team"] = proj["team_id"].map(abbr).fillna("FA")
    proj["gp"] = proj["avail"] * 82
    proj["injury_group"] = proj["player_id"].map(inj["group"] if len(inj) else {}).fillna("")
    mark = {"Injury-hit 2025-26": " ‡", "Healthy 2025-26, wrecked year before": " †", "": ""}
    proj["player"] = proj["name"] + proj["injury_group"].map(mark)
    chg = changes_vs(proj, compare, conn, year).set_index("player_id") if compare else None
    team_name = pd.read_sql_query("SELECT team_id, name FROM team", conn).set_index("team_id")["name"]

    per_game = pd.DataFrame({
        "Rank": proj["rank"], "Player": proj["player"], "Pos": proj["elig"], "Age": proj["age_t"].round(0), "Team": proj["team"], "Source": proj["source"],
        "Games out at start": proj["games_out"].round(0), "Proj GP (mean)": proj["gp"].round(0), "Typical GP (median)": proj["gp_median"].round(0),
        "GP P25": proj["gp_p25"].round(0), "GP P75": proj["gp_p75"].round(0), "MPG": proj["mpg"].round(1), "PPG": proj["pts"].round(1),
        "FGM": proj["fgm"].round(1), "FGA": proj["fga"].round(1), "FG%": (proj["fg_pct"] * 100).round(1), "3PM": proj["tpm"].round(1), "3PA": proj["tpa"].round(1),
        "FTM": proj["ftm"].round(1), "FTA": proj["fta"].round(1), "FT%": (proj["ft_pct"] * 100).round(1), "REB": proj["reb"].round(1), "AST": proj["ast"].round(1),
        "STL": proj["stl"].round(1), "BLK": proj["blk"].round(1), "TOV": proj["tov"].round(1), "Value (z total)": proj["z_total"].round(2),
        "Per-game rank": proj["pg_rank"], **{f"z {CAT_LABELS[c]}": proj[f"z_{c}"].round(2) for c in CATEGORIES},
    })
    changes_sheet = None
    if chg is not None:
        c = chg.reindex(proj["player_id"]).reset_index(drop=True)
        for label, col in (("Prev rank", "prev_rank"), ("Δ rank (+ = up)", "d_rank"), ("Δ value (z)", "d_z"), ("Prev PPG", "prev_pts"), ("Δ PPG", "d_pts"), ("Prev MPG", "prev_mpg"), ("Δ MPG", "d_mpg"),
                           ("Δ FGA", "d_fga"), ("Δ 3PM", "d_tpm"), ("Δ REB", "d_reb"), ("Δ AST", "d_ast"), ("Prev GP", "prev_gp"), ("Δ GP", "d_gp")):
            per_game[label] = c[col].round(0 if col in ("prev_rank", "d_rank", "prev_gp", "d_gp") else 1)
        big = proj[["player_id", "name", "team", "elig", "team_id", "rank", "pts", "mpg", "gp", "fga", "ast", "reb"]].reset_index(drop=True).join(c.drop(columns=["player_id"], errors="ignore"))
        big["z"] = proj["z_total"].reset_index(drop=True)
        big["from_team"] = big["last_team_id"].map(team_name)
        big = big[big["any_big"]].copy()
        big["Moved?"] = np.where(big["moved"], big["from_team"].fillna("") + " → " + big["team_id"].map(team_name).fillna("FA"), "")
        fmt = lambda a, b, d=1: [f"{x:.{d}f} → {y:.{d}f}" for x, y in zip(a, b)]
        big = big.assign(_k1=big["d_z"].abs(), _k2=big["d_pts"].abs()).sort_values(["_k1", "_k2"], ascending=False)
        changes_sheet = pd.DataFrame({
            "↕": np.where(big["d_z"] >= BIG["z"], "▲", np.where(big["d_z"] <= -BIG["z"], "▼", "")), "Player": big["name"], "Team": big["team"], "Pos": big["elig"],
            "Rank": big["rank"], "Prev rank": big["prev_rank"].round(0), "Δ rank (+ = up)": big["d_rank"].round(0), "Δ value (z)": big["d_z"].round(2), "Moved teams?": big["Moved?"],
            "PPG (prev → now)": fmt(big["prev_pts"], big["pts"]), "MPG (prev → now)": fmt(big["prev_mpg"], big["mpg"]), "GP (prev → now)": fmt(big["prev_gp"], big["gp"], 0),
            "Δ PPG": big["d_pts"].round(1), "Δ MPG": big["d_mpg"].round(1), "Δ FGA": big["d_fga"].round(1), "Δ AST": big["d_ast"].round(1),
            "Δ REB": big["d_reb"].round(1), "Δ 3PM": big["d_tpm"].round(1), "Δ GP": big["d_gp"].round(0),
        }).reset_index(drop=True)
    totals = pd.DataFrame({"Rank": proj["rank"], "Player": proj["player"], "Team": proj["team"], "Proj GP (mean)": proj["gp"].round(0)})
    for label, col in (("PTS", "pts"), ("FGM", "fgm"), ("FGA", "fga"), ("3PM", "tpm"), ("3PA", "tpa"), ("FTM", "ftm"), ("FTA", "fta"), ("REB", "reb"),
                       ("AST", "ast"), ("STL", "stl"), ("BLK", "blk"), ("TOV", "tov")):
        totals[label] = (proj[col] * proj["gp"]).round(0)
    totals["MIN"] = (proj["mpg"] * proj["gp"]).round(0)

    ka = proj[proj["games_out"] > 0]
    esp = ROOT / "data/external/known_absences.csv"
    man = ROOT / "data/external/known_absences_manual.csv"
    espn = pd.read_csv(esp).set_index("player_id") if esp.exists() else pd.DataFrame()
    manual = pd.read_csv(man) if man.exists() else pd.DataFrame(columns=["name", "games_out", "note"])
    manual_names = set(manual["name"].str.lower())
    absences = pd.DataFrame({
        "Rank": ka["rank"], "Player": ka["name"], "Team": ka["team"], "Games out at start": ka["games_out"].round(0),
        "Proj GP now": ka["gp"].round(0), "Proj GP if healthy at tip-off": (ka["avail_active"] * 82).round(0),
        "Source": ["manual" if n.lower() in manual_names else "ESPN" for n in ka["name"]],
        "ESPN est. return": [espn["est_return"].get(pid, "") if len(espn) else "" for pid in ka["player_id"]],
    })
    injured = proj[proj["injury_group"] != ""].copy()
    injured["gp_hist"] = injured["player_id"].map(inj["gp_hist"]); injured["rank_if_healthy"] = injured["player_id"].map(inj["rank_if_healthy"])
    injured = pd.DataFrame({
        "Group": injured["injury_group"], "Rank": injured["rank"], "Player": injured["name"], "Age": injured["age_t"].round(0),
        "GP 22-23 / 23-24 / 24-25 / 25-26": injured["gp_hist"], "Proj GP (mean)": injured["gp"].round(0), "Typical GP (median)": injured["gp_median"].round(0),
        "MPG": injured["mpg"].round(1), "Per-game rank": injured["pg_rank"], "Rank if 82% healthy": injured["rank_if_healthy"],
    }).sort_values(["Group", "Per-game rank"])

    top150 = proj[proj["rank"] <= 150]

    def row(label, d):
        return (label, len(d), round(d["gp"].mean(), 1), round(d["gp_median"].mean(), 1))
    means = pd.DataFrame([
        row("All projected players", proj), row("Projected MPG >= 15 and GP >= 10", proj[(proj.mpg >= 15) & (proj.gp >= 10)]),
        row("Projected starters (MPG >= 25)", proj[proj.mpg >= 25]), row("Top 100 on the board", proj.head(100)), row("Top 150 on the board", top150),
        row("  Top 150, not injury-impacted", top150[top150.injury_group == ""]), row("  Top 150, injury-impacted", top150[top150.injury_group != ""]),
    ], columns=["Group", "Players", "Mean GP", "Typical (median) GP"])
    notes = pd.DataFrame({"Notes": [
        f"Projections for the {year}-{str(year + 1)[-2:]} NBA season, in-house model (see docs/REPORT.md, docs/HANDOFF.md). Rank = board rank: nine-category value with counting stats scaled by projected games.",
        "Proj GP (mean) = expected games of 82: in-season availability for a player healthy at tip-off, reduced by 'Games out at start' (players already known to be out: ESPN's injuries page, status Out, plus hand-set overrides). The draft engine uses the mean.",
        "Typical GP (median) and P25 / P75 = what players like him actually played in the 2021-25 backtest at the same projected mean; the median is higher than the mean because season-ending injuries skew the average down. Display only.",
        "Per-game stats are averages when the player plays. Season totals = per game x Proj GP (mean). FG% / FT% are shown as percentages; z columns are category z-scores (TOV sign flipped) against the top-156 pool, scaled by projected games. Value (z total) is their sum with each category's penalty floored at -3 (a lost category is lost whether by 3 or 6 standard deviations; only Giannis-type free-throw/3PM outliers are affected), so for those players it exceeds the plain sum of the z columns.",
        "Source: 'model' = veterans projected from their last three seasons; 'rookie_tier' = rookies / no NBA history, by draft slot (all players in a tier share one profile, e.g. the five top-5 rookies); 'override' = hand-set. Position eligibility is approximated from nba_api labels, not ESPN's exact values.",
        "Injury group (‡ injury-hit 2025-26 / † healthy last year after a wrecked one) is inferred from games missed, not an injury feed. 'Rank if 82% healthy' = board rank if only that player's availability were raised to 82%.",
        "Players with no 2025-26 season (Haliburton, Lillard, Kyrie, VanVleet): availability is a 50/50 blend of the model and a shrunken baseline, since the model tends to under-project players returning from a lost season.",
    ]})

    if chg is not None:
        notes = pd.concat([notes, pd.DataFrame({"Notes": [
            f"HIGHLIGHTS compare this projection with the one {compare_label}. Green = higher / moved up the board, red = lower / moved down, for changes of at least: fantasy value (z total) {BIG['z']}, rank {BIG['rank']}, "
            f"PPG {BIG['pts']}, MPG {BIG['mpg']}, games {BIG['gp']:.0f}, FGA {BIG['fga']}, AST/REB {BIG['ast']}, 3PM {BIG['tpm']}; only players in the top 150 before or after are highlighted. "
            "A yellow Player cell means at least one of those is highlighted in his row. 'Biggest changes' lists them, largest change in fantasy value first (rank moves among players ranked 100+ are noisy: the board is tightly packed there), with the team he moved from where he changed teams. Δ columns at the right of Projections give every player's change."]})], ignore_index=True)

    out_path = Path(out); out_path.parent.mkdir(parents=True, exist_ok=True)
    from openpyxl.styles import Alignment, Font, PatternFill
    sheets = [("Projections", per_game)] + ([("Biggest changes", changes_sheet)] if changes_sheet is not None else []) + [
        ("Season totals", totals), ("Known absences", absences), ("Injury-impacted", injured), ("League GP means", means), ("Notes", notes)]
    fill = lambda hexcolor: PatternFill("solid", fgColor=hexcolor)
    with pd.ExcelWriter(out_path, engine="openpyxl") as xw:
        for name, df in sheets:
            df.to_excel(xw, sheet_name=name, index=False)
            ws = xw.book[name]
            for c in ws[1]:
                c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="DDE6F0"); c.alignment = Alignment(wrap_text=True, vertical="center")
            ws.freeze_panes = "C2" if name in ("Projections", "Biggest changes", "Season totals", "Known absences", "Injury-impacted") else "A2"
            if name not in ("Notes", "League GP means"):
                ws.auto_filter.ref = ws.dimensions
            for col in ws.columns:
                width = max(len(str(c.value)) if c.value is not None else 0 for c in col[:300])
                ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 7), 120 if name == "Notes" else 30)
            if name == "Notes":
                for r in ws.iter_rows(min_row=2):
                    r[0].alignment = Alignment(wrap_text=True, vertical="top")

        n_hl = 0
        if chg is not None:
            c = chg.reindex(proj["player_id"]).reset_index(drop=True)
            ws = xw.book["Projections"]
            hdr = {cell.value: cell.column for cell in ws[1]}
            cols = {"z": ("Value (z total)", "Δ value (z)"), "rank": ("Rank", "Δ rank (+ = up)"), "pts": ("PPG", "Δ PPG"), "mpg": ("MPG", "Δ MPG"), "gp": ("Proj GP (mean)", "Δ GP"),
                    "fga": ("FGA", "Δ FGA"), "ast": ("AST", "Δ AST"), "reb": ("REB", "Δ REB"), "tpm": ("3PM", "Δ 3PM")}
            for i in range(len(c)):
                if not c.at[i, "any_big"]:
                    continue
                ws.cell(row=i + 2, column=hdr["Player"]).fill = fill(YELLOW)
                for k, names in cols.items():
                    if c.at[i, f"big_{k}"]:
                        color = GREEN if c.at[i, f"d_{k}"] > 0 else RED
                        for nm in names:
                            ws.cell(row=i + 2, column=hdr[nm]).fill = fill(color)
                        n_hl += 1
            ws = xw.book["Biggest changes"]
            hdr = {cell.value: cell.column for cell in ws[1]}
            for r in range(2, ws.max_row + 1):
                for nm, thr in (("Δ value (z)", BIG["z"]), ("Δ rank (+ = up)", BIG["rank"]), ("Δ PPG", BIG["pts"]), ("Δ MPG", BIG["mpg"]), ("Δ GP", BIG["gp"]), ("Δ FGA", BIG["fga"]),
                                ("Δ AST", BIG["ast"]), ("Δ REB", BIG["reb"]), ("Δ 3PM", BIG["tpm"])):
                    cell = ws.cell(row=r, column=hdr[nm])
                    if isinstance(cell.value, (int, float)) and abs(cell.value) >= thr:
                        cell.fill = fill(GREEN if cell.value > 0 else RED)
                arrow = ws.cell(row=r, column=hdr["↕"])
                if arrow.value == "▲":
                    arrow.fill = fill(GREEN)
                elif arrow.value == "▼":
                    arrow.fill = fill(RED)
    extra = f", {int(chg['any_big'].sum())} players with big changes ({n_hl} highlighted values)" if chg is not None else ""
    print(f"wrote {out_path}  ({len(per_game)} players, {len(absences)} with known absences, {len(injured)} injury-impacted{extra})")


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--db-path", default=str(ROOT / "data/processed/nba.db"))
    ap.add_argument("--year", type=int, default=2026)
    ap.add_argument("--out", default=str(ROOT / "exports/player_projections_2026-27.xlsx"))
    ap.add_argument("--compare", default=None, help="another projections_<year>.parquet to diff against; the biggest changes are highlighted")
    ap.add_argument("--compare-label", default="before", help="how to describe that file in the notes, e.g. 'without team context'")
    a = ap.parse_args()
    main(a.db_path, a.year, a.out, a.compare, a.compare_label)
